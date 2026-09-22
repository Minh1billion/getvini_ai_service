import os

import pdfplumber
from docx import Document

SUPPORTED_EXTS = (".pdf", ".docx")

ROW_Y_TOLERANCE = 3.0
COLUMN_GAP_FACTOR = 2.2
COMPLEX_COLUMN_THRESHOLD = 2
COMPLEX_TABLE_THRESHOLD = 1


def _cluster_rows(words):
    rows = []
    for word in sorted(words, key=lambda w: (w["top"], w["x0"])):
        placed = False
        for row in rows:
            if abs(row["top"] - word["top"]) <= ROW_Y_TOLERANCE:
                row["words"].append(word)
                row["top"] = min(row["top"], word["top"])
                placed = True
                break
        if not placed:
            rows.append({"top": word["top"], "words": [word]})
    rows.sort(key=lambda r: r["top"])
    return rows


def _split_row_into_columns(row_words, char_width):
    row_words = sorted(row_words, key=lambda w: w["x0"])
    gap_threshold = max(char_width * COLUMN_GAP_FACTOR, 8.0)
    columns = []
    current = [row_words[0]]
    for prev, curr in zip(row_words, row_words[1:]):
        gap = curr["x0"] - prev["x1"]
        if gap > gap_threshold:
            columns.append(current)
            current = [curr]
        else:
            current.append(curr)
    columns.append(current)
    return columns


def _row_to_columns(row_words):
    if not row_words:
        return []
    char_widths = [w["x1"] - w["x0"] for w in row_words if w.get("text")]
    avg_char_width = (sum(char_widths) / len(char_widths)) if char_widths else 6.0
    column_groups = _split_row_into_columns(row_words, avg_char_width)
    columns = []
    for group in column_groups:
        group = sorted(group, key=lambda w: w["x0"])
        text = " ".join(w["text"] for w in group).strip()
        if text:
            columns.append(text)
    return columns


def _row_to_text(row_words):
    return "    ".join(_row_to_columns(row_words))


def _word_in_bboxes(word, bboxes):
    for x0, top, x1, bottom in bboxes:
        if word["x0"] >= x0 - 1 and word["x1"] <= x1 + 1 and word["top"] >= top - 1 and word["bottom"] <= bottom + 1:
            return True
    return False


def _extract_page_tables(page):
    blocks = []
    bboxes = []
    for table in page.find_tables():
        bboxes.append(table.bbox)
        data = table.extract() or []
        rows_text = []
        for row in data:
            cells = [(cell or "").strip() for cell in row]
            cells = [c for c in cells if c]
            if cells:
                rows_text.append(" | ".join(cells))
        if rows_text:
            blocks.append("\n".join(rows_text))
    return blocks, bboxes


def _detect_complex_layout(rows, tables_count):
    if tables_count >= COMPLEX_TABLE_THRESHOLD:
        return True
    for row in rows:
        if len(_row_to_columns(row["words"])) >= COMPLEX_COLUMN_THRESHOLD + 1:
            return True
    return False


def extract_text_from_pdf(path: str) -> tuple:
    parts = []
    has_complex_layout = False

    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            table_blocks, table_bboxes = _extract_page_tables(page)

            words = page.extract_words(use_text_flow=False, keep_blank_chars=False)
            non_table_words = [w for w in words if not _word_in_bboxes(w, table_bboxes)]
            rows = _cluster_rows(non_table_words)

            if _detect_complex_layout(rows, len(table_blocks)):
                has_complex_layout = True

            page_text = "\n".join(line for line in (_row_to_text(row["words"]) for row in rows) if line)

            page_parts = []
            if page_text:
                page_parts.append(page_text)
            page_parts.extend(table_blocks)

            if page_parts:
                parts.append("\n\n".join(page_parts))

    return "\n\n".join(parts).strip(), has_complex_layout


def extract_text_from_docx(path: str) -> tuple:
    document = Document(path)
    parts = []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [((cell.text or "").strip()) for cell in row.cells]
            cells = [c for c in cells if c]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip(), len(document.tables) >= COMPLEX_TABLE_THRESHOLD


def extract_text(path: str, filename: str = "") -> tuple:
    ext = os.path.splitext(filename or path)[1].lower()
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"Định dạng file không được hỗ trợ: {ext or 'không xác định'}, chỉ hỗ trợ .pdf và .docx")
    if ext == ".pdf":
        return extract_text_from_pdf(path)
    return extract_text_from_docx(path)
