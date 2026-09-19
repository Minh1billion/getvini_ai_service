import json
import os
import sys
import tempfile
import time

os.environ["QC_CACHE_DB"] = tempfile.mktemp(suffix=".sqlite3")
sys.path.insert(0, os.getcwd())

from fastapi.testclient import TestClient
from openpyxl import Workbook


def make_xlsx(path):
    wb = Workbook()
    a = wb.active
    a.title = "A"
    a["B1"] = "Danh sách sản phẩm"
    a["A2"] = "Kịch bản 1"
    a["D2"] = "Kịch bản Fix"
    for r, row in enumerate([["Sản phẩm sữa tươi", "giá 100k", "dung tích 500ml"], ["Bánh mì thịt", "giá 20k", "nặng 200g"], ["Trà xanh", "giá 15k", "chai 350ml"]], start=3):
        for c, v in enumerate(row, start=1):
            a.cell(r, c, v)
    a["D3"] = "Sản phẩm sữa chua"
    a["E3"] = "giá 12k"
    a["D4"] = "Trà đào"
    a["E4"] = "giá 18k"
    a["A7"] = "Kịch bản 1"
    a["A8"] = "Nước ngọtt vị chanh"
    a["B8"] = "https://example.com"
    a["A9"] = "Hello wrold this is a test"
    a["A11"] = "kịch  bản 3"
    b = wb.create_sheet("B")
    b["A1"] = "Xin chaoo cac ban"
    b["A2"] = "Product namee 123 abc@x.com"
    b["B3"] = "sản phẩm mới, giá 5.000đ"
    c = wb.create_sheet("C")
    c["A2"] = "Kịch bản X"
    c["A3"] = "aaa bbb"
    c["A4"] = "ccc ddd"
    c["A5"] = "Kịch bản Y"
    c["A6"] = "eee fff"
    c["C2"] = "Kịch bản Z"
    c["C3"] = "zzz"
    c["D3"] = "yyy"
    c["C4"] = "xxx"
    c["C5"] = "www"
    c["C6"] = "vvv"
    d = wb.create_sheet("D")
    d["C3"] = "Kịch bản Offset"
    d["C4"] = "dữ liệu lệch"
    d["D4"] = "giá 9k"
    wb.save(path)


def patch_llm():
    try:
        import app.infra.llm.client as m
    except ImportError:
        import app.feats.qc.llm_client as m

    class Msg:
        def __init__(self, content):
            self.content = content

    class Choice:
        def __init__(self, content):
            self.message = Msg(content)

    class Resp:
        def __init__(self, content):
            self.choices = [Choice(content)]

    class Completions:
        def create(self, **kw):
            payload = json.loads(kw["messages"][1]["content"])
            out = []
            for b in payload["content_blocks"]:
                out.append({
                    "id": b["id"],
                    "sheet": b["sheet"],
                    "scenario": b["scenario"],
                    "row_range": b["row_range"],
                    "product_ref": "P",
                    "attribute": "giá",
                    "claimed_value": "1",
                    "expected_value": "2",
                    "status": "mismatch",
                    "reasoning": "x",
                })
            out.append({"id": "zz", "sheet": "A", "scenario": "?", "row_range": [999, 1000], "status": "unresolved"})
            return Resp(json.dumps({"mismatches": out}, ensure_ascii=False))

    class Chat:
        completions = Completions()

    class FakeOpenAI:
        def __init__(self, **kw):
            self.chat = Chat()

    m.OpenAI = FakeOpenAI


def norm(obj, ids):
    s = json.dumps(obj, ensure_ascii=False, sort_keys=True)
    for i in ids:
        s = s.replace(i, "JOB")
    return json.loads(s)


def main(out_path):
    patch_llm()
    from app.main import app

    tmp = tempfile.mkdtemp()
    xlsx = os.path.join(tmp, "s.xlsx")
    make_xlsx(xlsx)
    txt = os.path.join(tmp, "t.txt")
    with open(txt, "w", encoding="utf-8") as f:
        f.write("Xin chaoo cac ban\n\nHello wrold\ntôi tên là Minhh\n")

    res = {}
    ids = []
    with TestClient(app) as cl:
        def post(name, url, **kw):
            r = cl.post(url, **kw)
            try:
                body = r.json()
            except Exception:
                body = r.text
            res[name] = {"status": r.status_code, "body": body}
            return r

        def files(path, name="file"):
            return {name: (os.path.basename(path), open(path, "rb"))}

        res["health"] = cl.get("/health").json()
        res["root"] = cl.get("/").json()
        post("check_sheets_xlsx", "/check/sheets", files=files(xlsx))
        post("check_sheets_txt", "/check/sheets", files=files(txt))
        post("check_text", "/check/text", data={"text": "Xin chaoo cac ban hello wrold https://a.com 123", "lang": "vi", "whitelist": "chaoo"})
        post("check_text_both", "/check/text", data={"text": "Xin chaoo cac ban hello wrold", "lang": "both"})
        post("check_text_en", "/check/text", data={"text": "Xin chaoo cac ban hello wrold", "lang": "en"})

        cases = {
            "xlsx_default": (xlsx, {}),
            "xlsx_all": (xlsx, {"sheet_names": "A,B,C,D", "lang": "both", "whitelist": "wrold, Nước"}),
            "xlsx_en": (xlsx, {"sheet_names": "B", "lang": "en"}),
            "txt": (txt, {"lang": "vi", "whitelist": "Minhh"}),
        }
        for name, (path, data) in cases.items():
            r = post(f"start_{name}", "/check/start", files=files(path), data=data)
            job_id = r.json().get("job_id")
            if not job_id:
                continue
            ids.append(job_id)
            for _ in range(200):
                j = cl.get(f"/check/{job_id}").json()
                if j["status"] != "running":
                    break
                time.sleep(0.05)
            res[f"job_{name}"] = j
            p = cl.get(f"/check/{job_id}/pdf")
            res[f"pdf_{name}"] = {"status": p.status_code, "ctype": p.headers.get("content-type"), "magic": p.content[:4].decode("latin1"), "disp": p.headers.get("content-disposition")}
            d = cl.delete(f"/check/{job_id}")
            res[f"delete_{name}"] = {"status": d.status_code, "body": d.json()}
            res[f"get_after_delete_{name}"] = cl.get(f"/check/{job_id}").status_code

        post("start_bad_sheet", "/check/start", files=files(xlsx), data={"sheet_names": "NOPE"})
        res["get_missing"] = cl.get("/check/nope").status_code

        r = cl.post("/check/stream", files=files(xlsx), data={"sheet_names": "A,B", "lang": "both"})
        res["stream"] = {"status": r.status_code, "ctype": r.headers.get("content-type"), "text": r.text}
        for line in r.text.splitlines():
            if line.startswith("data: "):
                ids.append(json.loads(line[6:])["job_id"])

        post("qc_sheets", "/qc/sheets", files=files(xlsx))
        post("qc_sheets_none", "/qc/sheets")
        post("qc_sheets_url", "/qc/sheets", data={"url": "http://x"})
        info = json.dumps([{"productName": "P", "specs": [{"key": "giá", "value": "1"}]}])
        for name, data in {
            "qc_run_A": {"sheet_name": "A", "product_info": info, "api_key": "k"},
            "qc_run_C_sel": {"sheet_name": "C", "product_info": info, "api_key": "k", "scenario_ids": "C::R1C0"},
            "qc_run_C_ok": {"sheet_name": "C", "product_info": info, "api_key": "k", "scenario_ids": "C::R2C0, C::R2C2"},
            "qc_run_B_none": {"sheet_name": "B", "product_info": info, "api_key": "k"},
            "qc_run_D": {"sheet_name": "D", "product_info": info, "api_key": "k", "batch_size": "1"},
            "qc_run_bad_json": {"sheet_name": "A", "product_info": "{", "api_key": "k"},
            "qc_run_bad_provider": {"sheet_name": "A", "product_info": info, "provider": "zzz"},
            "qc_run_no_info": {"sheet_name": "A"},
        }.items():
            post(name, "/qc/run", files=files(xlsx), data=data)
        post("qc_run_no_file", "/qc/run", data={"sheet_name": "A", "product_info": info})
        post("qc_run_info_file", "/qc/run", files={**files(xlsx), "product_info_file": ("p.json", info.encode())}, data={"sheet_name": "A", "api_key": "k"})

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(norm(res, ids), f, ensure_ascii=False, indent=1, sort_keys=True)


if __name__ == "__main__":
    main(sys.argv[1])
